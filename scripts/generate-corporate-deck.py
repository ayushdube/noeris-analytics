#!/usr/bin/env python3
"""Generate Noeris Analytics corporate prospect deck (.pptx)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Brand tokens (from globals.css / site)
BG = RGBColor(0x0E, 0x0E, 0x0D)
SURFACE = RGBColor(0x16, 0x16, 0x15)
FG = RGBColor(0xEC, 0xEC, 0xEA)
MUTED = RGBColor(0xB8, 0xB8, 0xB4)
MUTED_DARK = RGBColor(0x85, 0x85, 0x82)
ACCENT = RGBColor(0xFF, 0x02, 0x34)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_DISPLAY = "Syne"
FONT_BODY = "IBM Plex Sans"
FONT_MONO = "Consolas"

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "Noeris-Analytics-Corporate-Deck.pptx"


def set_slide_bg(slide, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, left=Inches(0.65), top=Inches(0.52), height=Inches(0.22)) -> None:
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.025), height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_label(slide, text: str, top=Inches(0.48), left=Inches(0.72)) -> None:
    box = slide.shapes.add_textbox(left, top, Inches(11), Inches(0.35))
    tf = box.text_frame
    tf.text = text.upper()
    p = tf.paragraphs[0]
    p.font.name = FONT_MONO
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = MUTED
    p.font.letter_spacing = Pt(1.2)


def add_section_header(
    slide,
    label: str,
    title: str,
    title_top=Inches(1.05),
    title_size=Pt(32),
) -> None:
    add_accent_bar(slide)
    add_label(slide, label)
    add_title(slide, title, top=title_top, size=title_size)


def add_title(
    slide,
    text: str,
    top=Inches(0.85),
    size=Pt(34),
    width=Inches(11.5),
    color=FG,
) -> None:
    lines = text.split("\n")
    box = slide.shapes.add_textbox(Inches(0.65), top, width, Inches(max(1.4, 0.55 * len(lines))))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT_DISPLAY
        p.font.size = size
        p.font.bold = False
        p.font.color.rgb = color
        p.line_spacing = 1.05


def add_subtitle(
    slide,
    text: str,
    top=Inches(1.75),
    size=Pt(28),
    color=MUTED,
    width=Inches(11.5),
) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), top, width, Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = text
    p = tf.paragraphs[0]
    p.font.name = FONT_DISPLAY
    p.font.size = size
    p.font.color.rgb = color
    p.line_spacing = 1.05


def add_body(
    slide,
    lines: list[str],
    top=Inches(2.1),
    size=Pt(14),
    color=MUTED,
    bullet=False,
    width=Inches(11.5),
) -> None:
    box = slide.shapes.add_textbox(Inches(0.65), top, width, Inches(4.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT_BODY
        p.font.size = size
        p.font.color.rgb = color
        p.line_spacing = 1.45
        p.space_after = Pt(8)
        if bullet:
            p.level = 0


def add_footer(slide, text: str = "noeris.io  ·  hello@noeris.io") -> None:
    box = slide.shapes.add_textbox(Inches(0.65), Inches(7.05), Inches(11), Inches(0.3))
    tf = box.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.name = FONT_MONO
    p.font.size = Pt(8)
    p.font.color.rgb = MUTED_DARK


def add_logo_lockup(slide, left=Inches(0.65), top=Inches(0.55)) -> None:
    mark = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, Inches(0.36), Inches(0.36)
    )
    mark.fill.solid()
    mark.fill.fore_color.rgb = ACCENT
    mark.line.fill.background()
    mark.adjustments[0] = 0.15

    n_box = slide.shapes.add_textbox(left, top + Inches(0.04), Inches(0.36), Inches(0.32))
    n_tf = n_box.text_frame
    n_tf.text = "N"
    n_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np = n_tf.paragraphs[0]
    np.alignment = PP_ALIGN.CENTER
    np.font.name = FONT_DISPLAY
    np.font.size = Pt(14)
    np.font.bold = True
    np.font.color.rgb = WHITE

    name = slide.shapes.add_textbox(left + Inches(0.48), top, Inches(2.5), Inches(0.22))
    name_tf = name.text_frame
    name_tf.text = "Noeris"
    name_p = name_tf.paragraphs[0]
    name_p.font.name = FONT_BODY
    name_p.font.size = Pt(11)
    name_p.font.bold = True
    name_p.font.color.rgb = FG

    sub = slide.shapes.add_textbox(left + Inches(0.48), top + Inches(0.18), Inches(2.5), Inches(0.18))
    sub_tf = sub.text_frame
    sub_tf.text = "ANALYTICS"
    sub_p = sub_tf.paragraphs[0]
    sub_p.font.name = FONT_MONO
    sub_p.font.size = Pt(6.5)
    sub_p.font.color.rgb = MUTED_DARK


def add_metric_cards(slide, metrics: list[tuple[str, str]], top=Inches(2.2)) -> None:
    card_w = Inches(2.65)
    gap = Inches(0.25)
    left = Inches(0.65)
    for metric, label in metrics:
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, card_w, Inches(1.35)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE
        card.line.color.rgb = RGBColor(0x28, 0x28, 0x27)

        m_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.22), card_w - Inches(0.4), Inches(0.5))
        m_tf = m_box.text_frame
        m_tf.text = metric
        mp = m_tf.paragraphs[0]
        mp.font.name = FONT_DISPLAY
        mp.font.size = Pt(28)
        mp.font.color.rgb = ACCENT

        l_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.72), card_w - Inches(0.4), Inches(0.5))
        l_tf = l_box.text_frame
        l_tf.word_wrap = True
        l_tf.text = label
        lp = l_tf.paragraphs[0]
        lp.font.name = FONT_BODY
        lp.font.size = Pt(10)
        lp.font.color.rgb = MUTED

        left += card_w + gap


def build_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — Cover
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_logo_lockup(s, top=Inches(0.7))
    add_label(s, "Boutique Data & AI Consultancy", top=Inches(1.35))
    add_title(
        s,
        "Enterprise Analytics Delivered by the\nConsultants Who Architect It.",
        top=Inches(1.85),
        size=Pt(40),
    )
    add_body(
        s,
        [
            "Turn fragmented business data into trusted insights for business teams and executive leadership.",
            "",
            "Principal-led engagements  ·  Power BI & Spotfire  ·  Cloud data platforms  ·  Agentic AI",
        ],
        top=Inches(4.0),
        size=Pt(13),
        color=MUTED,
    )
    add_footer(s)

    # 2 — Who we are
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(s, "About", "A boutique consultancy built for decision-makers.")
    add_body(
        s,
        [
            "Noeris Analytics is a principal-led data and AI consultancy. Every engagement is architected and delivered by senior consultants - not handed off to a delivery bench after the sale.",
            "",
            "Built by consultants who have delivered enterprise analytics solutions across North America, Europe, and APAC.",
        ],
        top=Inches(2.15),
        size=Pt(15),
    )
    add_metric_cards(
        s,
        [
            ("69", "Projects delivered"),
            ("4.9★", "Client rating"),
            ("40+", "Verified reviews"),
            ("Top", "Rated Plus"),
        ],
        top=Inches(4.55),
    )
    add_footer(s)

    # 3 — Problem
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_accent_bar(s)
    add_label(s, "The challenge")
    add_title(s, "You have the data.", top=Inches(1.05), size=Pt(36))
    add_subtitle(s, "But good data is only useful when you can act on it.", top=Inches(1.85), size=Pt(26))
    add_body(
        s,
        [
            "Most companies already own all the data they need. It lives in dozens of systems - none of them speaking the same language when leadership needs an answer.",
        ],
        top=Inches(2.75),
        size=Pt(14),
    )
    problems = [
        ("01", "Fragmented data across systems", "Sales, finance, operations, and marketing each maintain separate sources of truth."),
        ("02", "Reports that arrive too late", "Manual workflows delay visibility until the decision window has closed."),
        ("03", "Dashboards without decision context", "Charts without KPI hierarchy, ownership, or path from metric to action."),
        ("04", "Legacy analytics holding back growth", "On-prem warehouses and brittle ETL cannot scale with executive expectations."),
    ]
    y = Inches(3.55)
    for num, title, desc in problems:
        n_box = s.shapes.add_textbox(Inches(0.65), y, Inches(0.55), Inches(0.35))
        n_tf = n_box.text_frame
        n_tf.text = num
        np = n_tf.paragraphs[0]
        np.font.name = FONT_DISPLAY
        np.font.size = Pt(18)
        np.font.color.rgb = MUTED_DARK

        t_box = s.shapes.add_textbox(Inches(1.25), y, Inches(10.5), Inches(0.65))
        t_tf = t_box.text_frame
        t_tf.word_wrap = True
        t_tf.text = f"{title} — {desc}"
        tp = t_tf.paragraphs[0]
        tp.font.name = FONT_BODY
        tp.font.size = Pt(11)
        tp.font.color.rgb = MUTED
        y += Inches(0.72)
    add_footer(s)

    # 4 — Services
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(s, "Services", "End-to-end analytics capabilities.")
    services = [
        ("Business Intelligence", "Executive and operational dashboards aligned to how your business makes decisions.", "8–14 weeks"),
        ("Data Engineering", "Reliable pipelines and modern warehouses that make analytics trustworthy at scale.", "10–16 weeks"),
        ("Analytics Modernization", "Migrate legacy BI and reporting stacks without disrupting operations.", "12–20 weeks"),
        ("AI-powered Analytics", "Natural language access and intelligent summaries on trusted data foundations.", "6–12 weeks"),
        ("Agentic AI & Autonomous Analytics", "AI agents that query data, surface anomalies, and deliver narrative summaries.", "8–16 weeks"),
    ]
    y = Inches(2.0)
    for title, summary, timeline in services:
        bar = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), y + Inches(0.08), Inches(0.04), Inches(0.55))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()

        t_box = s.shapes.add_textbox(Inches(0.85), y, Inches(8.5), Inches(0.35))
        t_tf = t_box.text_frame
        t_tf.text = title
        tp = t_tf.paragraphs[0]
        tp.font.name = FONT_DISPLAY
        tp.font.size = Pt(15)
        tp.font.color.rgb = FG

        d_box = s.shapes.add_textbox(Inches(0.85), y + Inches(0.32), Inches(8.8), Inches(0.35))
        d_tf = d_box.text_frame
        d_tf.word_wrap = True
        d_tf.text = summary
        dp = d_tf.paragraphs[0]
        dp.font.name = FONT_BODY
        dp.font.size = Pt(10)
        dp.font.color.rgb = MUTED

        tm_box = s.shapes.add_textbox(Inches(10.2), y + Inches(0.05), Inches(2.5), Inches(0.3))
        tm_tf = tm_box.text_frame
        tm_tf.text = timeline
        tmp = tm_tf.paragraphs[0]
        tmp.font.name = FONT_MONO
        tmp.font.size = Pt(9)
        tmp.font.color.rgb = MUTED_DARK
        y += Inches(0.95)
    add_footer(s)

    # 5 — How we work
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(s, "Engagement model", "Principal-led from discovery to handover.")
    phases = [
        ("Principal-led Discovery", "Map decision workflows, data sources, and KPI gaps with leadership before touching architecture."),
        ("Principal-led Solution Architecture", "Design platform, semantic model, and integration patterns against security and compliance requirements."),
        ("Build", "Deliver pipelines, certified datasets, and dashboards in phased releases with parallel-run validation."),
        ("Adopt", "Train business teams, assign KPI ownership, and establish refresh monitoring."),
        ("Founder-led Executive Handover", "Transition with executive sign-off, documentation, and ongoing optimization ownership."),
    ]
    y = Inches(2.05)
    for i, (phase, detail) in enumerate(phases, 1):
        circle = s.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.65), y, Inches(0.38), Inches(0.38)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = SURFACE
        circle.line.color.rgb = ACCENT

        c_box = s.shapes.add_textbox(Inches(0.65), y + Inches(0.06), Inches(0.38), Inches(0.28))
        c_tf = c_box.text_frame
        c_tf.text = str(i)
        cp = c_tf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cp.font.name = FONT_MONO
        cp.font.size = Pt(11)
        cp.font.color.rgb = ACCENT

        t_box = s.shapes.add_textbox(Inches(1.2), y, Inches(10.5), Inches(0.65))
        t_tf = t_box.text_frame
        t_tf.word_wrap = True
        t_tf.text = phase
        tp = t_tf.paragraphs[0]
        tp.font.name = FONT_DISPLAY
        tp.font.size = Pt(14)
        tp.font.color.rgb = FG
        d = t_tf.add_paragraph()
        d.text = detail
        d.font.name = FONT_BODY
        d.font.size = Pt(10)
        d.font.color.rgb = MUTED
        d.line_spacing = 1.3
        y += Inches(0.95)
    add_footer(s)

    # 6 — Industries
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(
        s,
        "Industries",
        "Deep domain experience across regulated and high-growth sectors.",
        title_size=Pt(30),
    )
    industries = [
        "Oil & Gas",
        "Shopify & eCommerce",
        "Manufacturing",
        "SaaS",
        "Healthcare",
        "Financial Services",
        "Logistics",
        "Private Equity",
    ]
    cols = 2
    col_w = Inches(5.5)
    for i, ind in enumerate(industries):
        col = i % cols
        row = i // cols
        left = Inches(0.65) + col * (col_w + Inches(0.35))
        top = Inches(2.15) + row * Inches(0.72)

        pill = s.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, col_w, Inches(0.52)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = SURFACE
        pill.line.color.rgb = RGBColor(0x28, 0x28, 0x27)

        t_box = s.shapes.add_textbox(left + Inches(0.2), top + Inches(0.1), col_w - Inches(0.4), Inches(0.35))
        t_tf = t_box.text_frame
        t_tf.text = ind
        tp = t_tf.paragraphs[0]
        tp.font.name = FONT_BODY
        tp.font.size = Pt(13)
        tp.font.color.rgb = FG
    add_footer(s)

    # 7 — Outcomes
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(s, "Outcomes", "Measurable impact on how teams decide.")
    add_metric_cards(
        s,
        [
            ("40%", "Faster executive reporting cycles"),
            ("3x", "Increase in self-service adoption"),
            ("60%", "Reduction in manual reconciliation"),
            ("24h", "Data freshness for critical KPIs"),
        ],
        top=Inches(2.3),
    )
    add_body(
        s,
        [
            "Engagements are scoped around business outcomes - not tool implementations. KPI frameworks, certified datasets, and executive-ready reporting are delivered together.",
        ],
        top=Inches(4.1),
        size=Pt(14),
    )
    add_footer(s)

    # 8 — Case study 1
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(
        s,
        "Case study · Oil & Gas",
        "Production & Environmental Compliance Analytics",
        title_size=Pt(28),
    )
    add_body(
        s,
        [
            "Client: Integrated energy major",
            "",
            "Challenge: Production volumes and environmental compliance data lived in disconnected systems with no real-time monitoring - creating regulatory risk.",
            "",
            "Solution: All-in-One Production Reporting Tool and Air Monitoring System (EMS Stream) in Power BI, consolidating multi-site well data with real-time threshold alerting.",
            "",
            "Outcomes:",
            "• Real-time air quality monitoring across all well perimeters",
            "• Unified production tracking by well with automated refresh",
            "• Regulatory compliance reporting available on demand",
        ],
        top=Inches(2.0),
        size=Pt(12),
    )
    add_footer(s)

    # 9 — Case study 2
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(
        s,
        "Case study · Life Sciences",
        "Clinical Trial Operations Analytics Suite",
        title_size=Pt(28),
    )
    add_body(
        s,
        [
            "Client: Global contract research organisation",
            "",
            "Challenge: Analytics across trial operations, data quality, and HR sourced from fragmented CRM and HR systems with no single point of truth.",
            "",
            "Solution: Four interconnected Power BI reports plus full SAS Visual Analytics to Power BI migration.",
            "",
            "Outcomes:",
            "• Single analytics layer across clinical, HR, and data quality functions",
            "• Automated data credibility scoring replacing manual validation",
            "• Full SAS VA migration with no loss of functionality",
        ],
        top=Inches(2.0),
        size=Pt(12),
    )
    add_footer(s)

    # 10 — Technology
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(s, "Technology", "Platform-agnostic. Production-grade.")
    tech = [
        "Power BI", "TIBCO Spotfire", "Snowflake", "Databricks", "Azure", "AWS",
        "SQL Server", "PostgreSQL", "BigQuery", "Shopify", "HubSpot", "Salesforce",
        "NetSuite", "Dynamics 365", "SAP", "Python", "dbt",
    ]
    y = Inches(2.1)
    x_start = Inches(0.65)
    x = x_start
    max_x = Inches(12.5)
    for t in tech:
        pill_w = Inches(0.15 * len(t) + 0.55)
        if x + pill_w > max_x:
            x = x_start
            y += Inches(0.55)
        pill = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, pill_w, Inches(0.38))
        pill.fill.solid()
        pill.fill.fore_color.rgb = SURFACE
        pill.line.color.rgb = RGBColor(0x35, 0x35, 0x34)
        tb = s.shapes.add_textbox(x + Inches(0.12), y + Inches(0.06), pill_w - Inches(0.24), Inches(0.28))
        ttf = tb.text_frame
        ttf.text = t
        tp = ttf.paragraphs[0]
        tp.font.name = FONT_MONO
        tp.font.size = Pt(9)
        tp.font.color.rgb = MUTED
        x += pill_w + Inches(0.12)
    add_footer(s)

    # 11 — Team
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_section_header(s, "Who you work with", "Principal Consultants on every engagement.")
    founders = [
        (
            "Ankita Jaiswal",
            "Principal Analytics Consultant",
            "Business Intelligence · Executive Reporting · Power BI · TIBCO Spotfire",
        ),
        (
            "Ayush Dubey",
            "Principal Data Platform & AI Consultant",
            "Cloud Architecture · Data Engineering · Modern Analytics Platforms · AI",
        ),
    ]
    y = Inches(2.2)
    for name, title, focus in founders:
        card = s.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), y, Inches(11.8), Inches(1.35)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE
        card.line.color.rgb = RGBColor(0x28, 0x28, 0x27)

        n_box = s.shapes.add_textbox(Inches(0.9), y + Inches(0.22), Inches(5), Inches(0.35))
        n_tf = n_box.text_frame
        n_tf.text = name
        np = n_tf.paragraphs[0]
        np.font.name = FONT_DISPLAY
        np.font.size = Pt(18)
        np.font.color.rgb = FG

        t_box = s.shapes.add_textbox(Inches(0.9), y + Inches(0.55), Inches(10.5), Inches(0.3))
        t_tf = t_box.text_frame
        t_tf.text = title
        tp = t_tf.paragraphs[0]
        tp.font.name = FONT_BODY
        tp.font.size = Pt(11)
        tp.font.color.rgb = ACCENT

        f_box = s.shapes.add_textbox(Inches(0.9), y + Inches(0.85), Inches(10.5), Inches(0.35))
        f_tf = f_box.text_frame
        f_tf.text = focus
        fp = f_tf.paragraphs[0]
        fp.font.name = FONT_BODY
        fp.font.size = Pt(10)
        fp.font.color.rgb = MUTED
        y += Inches(1.55)

    add_body(
        s,
        [
            "Microsoft-Certified · Expert-Vetted (Top Rated Plus, Upwork) · Delivered across 3 continents · Principal-led on every engagement",
        ],
        top=Inches(5.35),
        size=Pt(11),
        color=MUTED_DARK,
    )
    add_footer(s)

    # 12 — Testimonial
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_accent_bar(s, top=Inches(1.15), height=Inches(0.22))
    add_label(s, "Client feedback", top=Inches(1.1), left=Inches(0.72))
    quote_box = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.65), Inches(11.8), Inches(2.8)
    )
    quote_box.fill.solid()
    quote_box.fill.fore_color.rgb = SURFACE
    quote_box.line.color.rgb = RGBColor(0x28, 0x28, 0x27)

    q = s.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.0), Inches(2.0))
    q_tf = q.text_frame
    q_tf.word_wrap = True
    q_tf.text = (
        '"Ankita is one of the most skilled professionals I have worked with. '
        'She has advanced Power BI knowledge along with outstanding communication skills."'
    )
    qp = q_tf.paragraphs[0]
    qp.font.name = FONT_DISPLAY
    qp.font.size = Pt(22)
    qp.font.color.rgb = FG
    qp.line_spacing = 1.35

    attr = s.shapes.add_textbox(Inches(1.0), Inches(4.15), Inches(10), Inches(0.35))
    attr_tf = attr.text_frame
    attr_tf.text = "Verified client (Upwork)"
    ap = attr_tf.paragraphs[0]
    ap.font.name = FONT_MONO
    ap.font.size = Pt(10)
    ap.font.color.rgb = MUTED_DARK
    add_footer(s)

    # 13 — CTA
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_logo_lockup(s, top=Inches(0.85))
    add_label(s, "Next steps", top=Inches(1.45), left=Inches(0.65))
    add_title(s, "Work directly with the consultants", top=Inches(1.85), size=Pt(36))
    add_subtitle(
        s,
        "who architect and deliver every engagement.",
        top=Inches(2.55),
        size=Pt(30),
        color=MUTED,
    )

    contact_card = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(3.55), Inches(6.2), Inches(1.55)
    )
    contact_card.fill.solid()
    contact_card.fill.fore_color.rgb = SURFACE
    contact_card.line.color.rgb = RGBColor(0x28, 0x28, 0x27)

    email_box = s.shapes.add_textbox(Inches(0.95), Inches(3.82), Inches(5.5), Inches(0.35))
    email_tf = email_box.text_frame
    email_tf.text = "hello@noeris.io"
    email_p = email_tf.paragraphs[0]
    email_p.font.name = FONT_MONO
    email_p.font.size = Pt(11)
    email_p.font.color.rgb = FG

    web_box = s.shapes.add_textbox(Inches(0.95), Inches(4.2), Inches(5.5), Inches(0.35))
    web_tf = web_box.text_frame
    web_tf.text = "noeris.io/contact"
    web_p = web_tf.paragraphs[0]
    web_p.font.name = FONT_MONO
    web_p.font.size = Pt(10)
    web_p.font.color.rgb = MUTED

    cta = s.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(5.45), Inches(3.45), Inches(0.58)
    )
    cta.fill.solid()
    cta.fill.fore_color.rgb = ACCENT
    cta.line.fill.background()
    cta_box = s.shapes.add_textbox(Inches(0.9), Inches(5.58), Inches(3.1), Inches(0.35))
    cta_tf = cta_box.text_frame
    cta_tf.text = "Book a call  →"
    cp = cta_tf.paragraphs[0]
    cp.font.name = FONT_BODY
    cp.font.size = Pt(13)
    cp.font.bold = True
    cp.font.color.rgb = WHITE
    add_footer(s, "Confidential · Noeris Analytics · noeris.io")

    return prs


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs = build_deck()
    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
